# -*- coding: utf-8 -*-
"""Arma la Figura 3.4 como diapositiva EDITABLE de PowerPoint.

POR QUE EXISTE ESTE SCRIPT
--------------------------
El SVG que dejan los demas generadores es editable en Inkscape, pero Inkscape es
una herramienta mas que aprender. En PowerPoint, en cambio, se toca el cuerpo de
letra de cualquier rotulo con dos clics. Este script arma la misma Figura 3.4 pero
como diapositiva: los dos mapas y la barra de color van como imagen -son rasteres,
no hay otra-, y TODO EL TEXTO es un cuadro de texto de PowerPoint, y cada color de
la leyenda es un rectangulo con relleno solido.

EL TAMANO NO ES CAPRICHOSO. La diapositiva mide 15,92 cm de ancho, que es
EXACTAMENTE el ancho que ocupa la figura dentro del .docx (el ancho util de una A4
con margenes de una pulgada). Por eso los cuerpos de letra que se ven aca son los
que se van a imprimir: 8 pt es 8 pt. Si se cambia el ancho de la diapositiva, esa
correspondencia se pierde.

COMO SE VUELVE A LLEVAR AL WORD
   PowerPoint -> Archivo -> Exportar -> Cambiar tipo de archivo -> PNG,
   y elegir "Solo esta diapositiva". Insertar ese PNG en el Word con un ancho de
   15,92 cm. El alto sale solo.

REQUISITOS   pip install python-pptx        (ademas de numpy, rasterio, matplotlib)
ENTRADA      02_Subsets_SNAP_QGIS/Sentinel_2/02_pre*|03_post*/BOSQUE_NW_02/*.tif
SALIDA       figuras/editables_pptx/Figura_3.4_dNBR_severidad_editable.pptx
             y los tres bitmaps auxiliares, en esa misma carpeta

USO   python gen_tp3_fig_dnbr_pptx.py
"""
import glob, os
import numpy as np, rasterio
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap, BoundaryNorm
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.abspath(os.path.join(AQUI, "..", "editables_pptx"))
PROY = os.path.abspath(os.path.join(AQUI, "..", "..", "..", ".."))
S2 = os.path.join(PROY, "TP3_Datos_Opticos", "02_Subsets_SNAP_QGIS", "Sentinel_2")
os.makedirs(SALIDA, exist_ok=True)

AOI = "BOSQUE_NW_02"
MAL = [0, 1, 3, 8, 9, 10, 11]
# Los mismos siete umbrales contiguos de TP3_04_dnbr_incendio.py, que manda.
CLASES = [("Regeneración alta", "#1a9850", -np.inf, -0.250),
          ("Regeneración baja", "#a6d96a", -0.250, -0.100),
          ("Sin cambio",        "#d9d9d9", -0.100,  0.100),
          ("Baja",              "#ffeda0",  0.100,  0.270),
          ("Moderada-baja",     "#feb24c",  0.270,  0.440),
          ("Moderada-alta",     "#f4622e",  0.440,  0.660),
          ("Alta",              "#a50f15",  0.660,  np.inf)]


def nbr(pat):
    f = sorted(glob.glob(os.path.join(S2, "*", AOI, pat)))[0]
    with rasterio.open(f) as d:
        ni, sw, scl = d.read(7).astype("f4"), d.read(10).astype("f4"), d.read(11)
    v = (~np.isin(scl, MAL)) & (ni > 0) & (sw > 0)
    return np.where(v, (ni - sw) / np.where(ni + sw == 0, 1, ni + sw), np.nan).astype("f4"), v


a, va = nbr("*20251125*.tif"); b, vb = nbr("*20260305*.tif")
v = va & vb; d = np.where(v, a - b, np.nan); tot = int(v.sum())
sev = np.full(d.shape, np.nan)
for i, (_, _, lo, hi) in enumerate(CLASES):
    sev[v & (d >= lo) & (d < hi)] = i


def bitmap(arr, cmap, norm, nombre):
    """Guarda el panel SIN ejes, SIN titulo y SIN leyenda: solo el raster."""
    fig = plt.figure(figsize=(6, 6), dpi=250); ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    if norm is None:
        ax.imshow(arr, cmap=cmap, vmin=-0.3, vmax=1.0)
    else:
        ax.imshow(arr, cmap=cmap, norm=norm)
    fig.savefig(os.path.join(SALIDA, nombre), dpi=250, pad_inches=0); plt.close(fig)


bitmap(d, "RdYlGn_r", None, "panel_dnbr.png")
cmap = ListedColormap([c for _, c, _, _ in CLASES])
bitmap(sev, cmap, BoundaryNorm(range(len(CLASES) + 1), cmap.N), "panel_severidad.png")
g = np.linspace(1.0, -0.3, 600).reshape(-1, 1)
fig = plt.figure(figsize=(0.5, 6), dpi=250); ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
ax.imshow(g, cmap="RdYlGn_r", vmin=-0.3, vmax=1.0, aspect="auto")
fig.savefig(os.path.join(SALIDA, "barra_dnbr.png"), dpi=250, pad_inches=0); plt.close(fig)

W, H = 15.92, 9.6
TIPO = "Times New Roman"
CLASES = [("Alta", "A50F15", "9.668,7 ha (43,1 %)"),
          ("Moderada-alta", "F4622E", "3.672,0 ha (16,4 %)"),
          ("Moderada-baja", "FEB24C", "2.825,6 ha (12,6 %)"),
          ("Baja", "FFEDA0", "1.842,8 ha (8,2 %)"),
          ("Sin cambio", "D9D9D9", "4.170,7 ha (18,6 %)"),
          ("Regeneración baja", "A6D96A", "257,0 ha (1,1 %)"),
          ("Regeneración alta", "1A9850", "8,3 ha (0,0 %)")]

prs = Presentation(); prs.slide_width = Cm(W); prs.slide_height = Cm(H)
s = prs.slides.add_slide(prs.slide_layouts[6])

def texto(x, y, w, h, t, pt=8, negrita=False, al=PP_ALIGN.CENTER, color="000000"):
    c = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = c.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, linea in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        r = p.add_run(); r.text = linea
        r.font.size = Pt(pt); r.font.bold = negrita
        r.font.name = TIPO; r.font.color.rgb = RGBColor.from_string(color)
    return c

# --- titulo
texto(0.2, 0.15, W - 0.4, 0.95,
      "El incendio medido con dNBR — AOI de bosque\n"
      "pre: 25/11/2025 (escena limpia)  |  post: 05/03/2026", pt=9, negrita=True)

# --- paneles
PY, PL = 1.62, 6.30
texto(0.35, 1.15, PL, 0.42, "dNBR (continuo)", pt=8.5, negrita=True)
texto(9.27, 1.15, PL, 0.42, "Severidad (Key y Benson, siete clases)", pt=8.5, negrita=True)
s.shapes.add_picture(os.path.join(SALIDA, "panel_dnbr.png"), Cm(0.35), Cm(PY), Cm(PL), Cm(PL))
s.shapes.add_picture(os.path.join(SALIDA, "panel_severidad.png"), Cm(9.27), Cm(PY), Cm(PL), Cm(PL))

# --- barra de color con sus numeros como texto
BX, BW = 7.05, 0.34
s.shapes.add_picture(os.path.join(SALIDA, "barra_dnbr.png"), Cm(BX), Cm(PY), Cm(BW), Cm(PL))
for val in (1.0, 0.8, 0.6, 0.4, 0.2, 0.0, -0.2):
    y = PY + (1.0 - val) / 1.3 * PL
    texto(BX + BW + 0.06, y - 0.16, 0.85, 0.32,
          ("%.1f" % val).replace(".", ","), pt=7, al=PP_ALIGN.LEFT)
texto(BX - 0.45, PY - 0.40, 1.6, 0.34, "dNBR", pt=7, al=PP_ALIGN.CENTER)

# --- leyenda: FUERA de los mapas, al pie. Dos columnas por cuatro filas.
LY, FILA, COL = 8.10, 0.34, 7.6
for i, (nom, hexcol, dato) in enumerate(CLASES):
    fila, col = i % 4, i // 4
    x = 0.35 + col * COL; y = LY + fila * FILA
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y + 0.06), Cm(0.42), Cm(0.20))
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor.from_string(hexcol)
    r.line.color.rgb = RGBColor.from_string("595959"); r.line.width = Pt(0.5)
    r.shadow.inherit = False
    texto(x + 0.55, y, COL - 0.9, 0.32, "%s — %s" % (nom, dato), pt=8, al=PP_ALIGN.LEFT)

prs.save(os.path.join(SALIDA, "Figura_3.4_dNBR_severidad_editable.pptx"))
print("listo -> %s" % os.path.join(SALIDA, "Figura_3.4_dNBR_severidad_editable.pptx"))
print("   diapositiva de %.2f x %.2f cm, todo el texto editable" % (W, H))
