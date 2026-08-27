# -*- coding: utf-8 -*-
"""Convierte una figura de matplotlib en una diapositiva EDITABLE de PowerPoint.

QUE PROBLEMA RESUELVE
---------------------
Las figuras de este proyecto se dibujan con matplotlib a 15-17 pulgadas de ancho
(unos 40 cm) y despues se insertan en el .docx a 15,92 cm. Eso encoge todo por un
factor de 0,4: un rotulo declarado a 12 pt se IMPRIME a unos 5 pt. El SVG permite
retocarlo, pero exige Inkscape.

Este modulo deja la misma figura como diapositiva de PowerPoint: el dibujo va como
imagen de fondo -las lineas, los puntos y los mapas son rasteres o vectores que no
hace falta tocar- y CADA ROTULO se convierte en un cuadro de texto de PowerPoint,
colocado en su posicion exacta y con el cuerpo de letra REAL que va a imprimirse.
Abrir el .pptx y agrandar un rotulo es cuestion de dos clics.

La diapositiva mide 15,92 cm de ancho, que es el ancho util de la A4 con margenes
de una pulgada: el mismo con el que la figura entra en el Word. Por eso los puntos
son puntos de verdad.

USO
    from _a_pptx import a_pptx
    a_pptx(fig, "tp3_fig_saturacion", CARPETA_DE_SALIDA)

REQUISITO   pip install python-pptx
"""
import os
import matplotlib.text as mtext
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ANCHO_CM = 15.92          # ancho util de la A4 con margenes de 1 pulgada
TIPO = "Times New Roman"


# matplotlib entiende mathtext ($...$); PowerPoint no. Se traduce lo poco que
# aparece en las figuras del proyecto y, si queda algo, se le sacan los signos $.
_MATH = {r"\rightarrow": "→", r"\to": "→", r"\leftarrow": "←",
         r"\gamma": "γ", r"\sigma": "σ", r"\alpha": "α", r"\lambda": "λ",
         r"\times": "×", r"\leq": "≤", r"\geq": "≥", r"\approx": "≈",
         r"^0": "⁰", r"^2": "²", "\\,": " ", "{": "", "}": ""}


def _sin_mathtext(t):
    if "$" not in t:
        return t
    import re as _re
    def _tr(m):
        x = m.group(1)
        for k, v in _MATH.items():
            x = x.replace(k, v)
        return x
    return _re.sub(r"\$([^$]*)\$", _tr, t)


def _color(t):
    c = t.get_color()
    from matplotlib.colors import to_hex
    return to_hex(c, keep_alpha=False)[1:].upper()


def a_pptx(fig, nombre, salida, ancho_cm=ANCHO_CM, dpi=250):
    os.makedirs(salida, exist_ok=True)
    fig.canvas.draw()
    r = fig.canvas.get_renderer()
    W_px, H_px = fig.canvas.get_width_height()
    W_in = fig.get_figwidth()
    alto_cm = ancho_cm * H_px / float(W_px)
    # los pt de matplotlib estan referidos al ancho en pulgadas de la figura;
    # al llevarla a ancho_cm hay que escalarlos por la misma razon
    k = (ancho_cm / 2.54) / W_in

    # Los rotulos de marcas que caen FUERA de los limites del eje siguen
    # existiendo como objetos Text aunque no se dibujen. Si no se los excluye,
    # aparecen en el .pptx como numeros sueltos flotando fuera del grafico.
    fuera = set()
    for ax in fig.axes:
        for eje, lims, ticks, labs in (
                ("x", ax.get_xlim(), ax.get_xticks(), ax.get_xticklabels()),
                ("y", ax.get_ylim(), ax.get_yticks(), ax.get_yticklabels())):
            lo, hi = min(lims), max(lims)
            for val, lab in zip(ticks, labs):
                if not (lo - 1e-9 <= val <= hi + 1e-9):
                    fuera.add(id(lab))

    # Un rotulo que en la figura quedaba TAPADO por el recuadro opaco de una
    # leyenda, en el .pptx quedaria por encima (el fondo es una sola imagen).
    # Se excluyen los que caen dentro del recuadro de una leyenda sin ser suyos.
    marcos, propios = [], set()
    for ax in fig.axes:
        lg = ax.get_legend()
        if lg is None:
            continue
        try:
            marcos.append(lg.get_window_extent(renderer=r))
        except Exception:
            continue
        for t in lg.findobj(mtext.Text):
            propios.add(id(t))

    def tapado(t, bb):
        if id(t) in propios:
            return False
        xc, yc = (bb.x0 + bb.x1) / 2.0, (bb.y0 + bb.y1) / 2.0
        return any(m.x0 <= xc <= m.x1 and m.y0 <= yc <= m.y1 for m in marcos)

    textos = []
    for t in fig.findobj(mtext.Text):
        if id(t) in fuera:
            continue
        if not t.get_visible() or not (t.get_text() or "").strip():
            continue
        try:
            bb = t.get_window_extent(renderer=r)
        except Exception:
            continue
        if bb.width <= 0 or bb.height <= 0 or tapado(t, bb):
            continue
        textos.append((t, bb, t.get_fontsize() * k, t.get_rotation(),
                       t.get_fontweight() in ("bold", "heavy", "black", 600, 700, 800, 900),
                       _color(t)))

    # --- fondo: la MISMA figura pero sin una sola letra
    for t, *_ in textos:
        t.set_visible(False)
    fondo = os.path.join(salida, nombre + "_fondo.png")
    fig.savefig(fondo, dpi=dpi)          # SIN bbox_inches="tight": las coordenadas
    for t, *_ in textos:                 # tienen que seguir siendo las del lienzo
        t.set_visible(True)

    # --- diapositiva
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(ancho_cm), Cm(alto_cm)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(fondo, 0, 0, Cm(ancho_cm), Cm(alto_cm))

    cx, cy = ancho_cm / float(W_px), alto_cm / float(H_px)
    for t, bb, pt, rot, negrita, col in textos:
        an, al = bb.width * cx, bb.height * cy
        xc, yc = (bb.x0 + bb.x1) / 2.0 * cx, alto_cm - (bb.y0 + bb.y1) / 2.0 * cy
        if abs(rot) > 1:                       # texto girado: se intercambian los lados
            an, al = al, an
        an += 0.30; al += 0.12                 # holgura para que no se recorte al agrandar
        c = s.shapes.add_textbox(Cm(xc - an / 2.0), Cm(yc - al / 2.0), Cm(an), Cm(al))
        tf = c.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, linea in enumerate(_sin_mathtext(t.get_text() or "").split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run(); run.text = linea
            run.font.size = Pt(max(1, round(pt * 2) / 2.0))
            run.font.bold = negrita
            run.font.name = TIPO
            try:
                run.font.color.rgb = RGBColor.from_string(col)
            except Exception:
                pass
        if abs(rot) > 1:
            c.rotation = -rot % 360

    ruta = os.path.join(salida, nombre + "_editable.pptx")
    prs.save(ruta)
    ptmin = min(x[2] for x in textos) if textos else 0
    ptmax = max(x[2] for x in textos) if textos else 0
    print("  -> %s_editable.pptx  (%.2f x %.2f cm, %d rotulos, de %.1f a %.1f pt reales)"
          % (nombre, ancho_cm, alto_cm, len(textos), ptmin, ptmax))
    return ruta
